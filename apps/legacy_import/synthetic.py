"""Synthetic schema fixture only. Never used as a real golden master."""

from datetime import datetime, timedelta

import openpyxl

from apps.legacy_import.schema import CHECKS, DISTRIBUTION


def generate_synthetic_legacy_workbook(
    filepath, total=282, g1_count=9, g2_count=34, g3_count=212, sem_grupo_count=27
):
    assert total == g1_count + g2_count + g3_count + sem_grupo_count
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DISTRIBUIÇÃO"
    for key, col in DISTRIBUTION.items():
        ws[f"{col}2"] = key
    classification = wb.create_sheet("CLASSIFICAÇÃO")
    classification["Z3"] = "RECIFE"
    analyst = wb.create_sheet("ANÁLISE - TESTE")
    wb.create_sheet("REVISÃO")
    wb.create_sheet("DILIGÊNCIA")
    wb.create_sheet("VERIF. CNPJs")
    for i in range(total):
        base = f"{i + 1:08d}0001"
        digits = [int(c) for c in base]
        for weights in (
            [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2],
            [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2],
        ):
            rem = sum(d * w for d, w in zip(digits, weights, strict=True)) % 11
            digits.append(0 if rem < 2 else 11 - rem)
        cnpj = "".join(str(d) for d in digits)
        dt = datetime(2025, 9, 19, 9) + timedelta(minutes=i)
        group = (
            "G1"
            if i < g1_count
            else "G2"
            if i < g1_count + g2_count
            else "G3"
            if i < g1_count + g2_count + g3_count
            else "SEM_GRUPO"
        )
        values = {
            "number": i + 1,
            "institution": f"Synthetic institution {i + 1}",
            "cnpj": cnpj,
            "processo_sei": f"71000.{i + 1:06d}/2025-01",
            "date": dt.date(),
            "time": dt.time(),
            "group": group,
            "analyst": "TESTE",
            "municipality": "Recife" if group == "G2" else "Campinas",
            "state": "PE" if group == "G2" else "SP",
            "vagas_femininas": 10 if group == "G1" else 0,
            "vagas_masculinas": 10 if group in ("G2", "G3") else 0,
            "vagas_maes_nutrizes": 0,
            "vagas_solicitadas": 0 if group == "SEM_GRUPO" else 10,
            "capacidade_total": 20,
        }
        for key, col in DISTRIBUTION.items():
            ws[f"{col}{i + 3}"] = values.get(key)
        r = i + 4
        for col, key in zip(
            "ABCD", ["institution", "cnpj", "processo_sei", "analyst"], strict=True
        ):
            analyst[f"{col}{r}"] = values[key]
        for columns, _evidence in CHECKS.values():
            for col in columns:
                analyst[f"{col}{r}"] = "ISENTO" if col == "BV" else "ATENDE"
        analyst[f"BY{r}"] = "APTA"
    wb.save(filepath)
    return filepath
