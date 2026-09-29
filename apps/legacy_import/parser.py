"""Read-only XLSX parser; raw values and formula caches remain independent observations."""

from dataclasses import dataclass, field
from hashlib import sha256
from pathlib import Path

import openpyxl

from . import schema
from .normalizers import json_value, text


@dataclass
class WorkbookData:
    filename: str
    sha256: str
    sheets: list
    epoch: object
    distribution: list = field(default_factory=list)
    analyses: list = field(default_factory=list)
    reviews: list = field(default_factory=list)
    diligences: list = field(default_factory=list)
    ranking: list = field(default_factory=list)
    pronasci: list = field(default_factory=list)
    metrics: dict = field(default_factory=dict)
    consolidated: list = field(default_factory=list)
    issues: list = field(default_factory=list)


def parse_workbook(path):
    path = Path(path)
    digest = sha256(path.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    values = openpyxl.load_workbook(path, data_only=True, read_only=False)
    data = WorkbookData(path.name, digest, values.sheetnames, values.epoch)

    def rows(sheet, mapping, start):
        if sheet not in values:
            data.issues.append(
                {
                    "severity": "ERROR",
                    "sheet": sheet,
                    "row": None,
                    "column": "",
                    "processo_sei": "",
                    "code": "MISSING_SHEET",
                    "message": "Required source sheet missing",
                    "raw_value": None,
                }
            )
            return []
        result = []
        ws, fs = values[sheet], formulas[sheet]
        for row in range(start, ws.max_row + 1):
            sei = text(ws[f"{mapping['processo_sei']}{row}"].value)
            identity = ws[f"{mapping.get('cnpj', mapping['processo_sei'])}{row}"].value
            if not sei and identity is None:
                continue
            if not sei:
                data.issues.append(
                    {
                        "severity": "ERROR",
                        "sheet": sheet,
                        "row": row,
                        "column": mapping["processo_sei"],
                        "processo_sei": "",
                        "code": "MISSING_SEI",
                        "message": "Identity row lacks process SEI; excluded from recognized universe",
                        "raw_value": json_value(identity),
                    }
                )
                continue
            cells = {col: json_value(ws[f"{col}{row}"].value) for col in mapping.values()}
            cached = {col: ws[f"{col}{row}"].value for col in mapping.values()}
            source_formulas = {}
            for col in mapping.values():
                cell = fs[f"{col}{row}"]
                if cell.data_type == "f" or hasattr(cell.value, "text"):
                    source_formulas[col] = json_value(cell.value)
                    if cached[col] is None:
                        data.issues.append(
                            {
                                "severity": "WARNING",
                                "sheet": sheet,
                                "row": row,
                                "column": col,
                                "processo_sei": sei,
                                "code": "MISSING_FORMULA_CACHE",
                                "message": "Formula has no cached value",
                                "raw_value": source_formulas[col],
                            }
                        )
            result.append(
                {
                    "sheet": sheet,
                    "row": row,
                    "processo_sei": sei,
                    "values": {key: cached[col] for key, col in mapping.items()},
                    "raw": cells,
                    "formulas": source_formulas,
                }
            )
        return result

    try:
        data.distribution = rows("DISTRIBUIÇÃO", schema.DISTRIBUTION, 3)
        analysis_map = {**schema.ANALYST_IDENTITY, **schema.ANALYST_RESULT}
        for cols, evidence in schema.CHECKS.values():
            for col in [*cols, *evidence.values()]:
                analysis_map[col] = col
        for sheet in values.sheetnames:
            if sheet.startswith("ANÁLISE -"):
                data.analyses.extend(rows(sheet, analysis_map, 4))
        data.reviews = rows("REVISÃO", schema.REVIEW, 4)
        data.diligences = rows("DILIGÊNCIA", schema.DILIGENCE, 4)
        if "ANÁLISE" in values:
            data.consolidated = rows(
                "ANÁLISE", {**schema.ANALYST_IDENTITY, **schema.ANALYST_RESULT}, 4
            )
        if "CLASSIFICAÇÃO" in values:
            for group, mapping in schema.RANKING.items():
                for record in rows("CLASSIFICAÇÃO", mapping, 3):
                    record["block"] = group
                    data.ranking.append(record)
            data.pronasci = [
                text(r[0])
                for r in values["CLASSIFICAÇÃO"].iter_rows(min_col=26, max_col=26, values_only=True)
                if r[0] is not None
            ]
        if "MÉTRICAS" in values:
            data.metrics = {
                cell.coordinate: json_value(cell.value)
                for row in values["MÉTRICAS"]
                for cell in row
                if cell.value is not None
            }
    finally:
        formulas.close()
        values.close()
    return data
