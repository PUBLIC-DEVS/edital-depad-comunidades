"""Synthetic malformed cases verify safety, never establish real parity."""

import openpyxl
import pytest
from django.core.exceptions import ValidationError

from apps.legacy_import.importer import LegacyImporter
from apps.legacy_import.models import LegacyImportIssue, LegacyImportRun
from apps.submissions.models import Submission
from apps.submissions.services.legacy_generator import generate_synthetic_legacy_workbook

pytestmark = [pytest.mark.legacy, pytest.mark.django_db]


def source(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    generate_synthetic_legacy_workbook(
        path, total=2, g1_count=0, g2_count=0, g3_count=2, sem_grupo_count=0
    )
    return path


def test_strict_invalid_timestamp_rolls_back_domain_and_keeps_issue(tmp_path):
    path = source(tmp_path)
    workbook = openpyxl.load_workbook(path)
    workbook["DISTRIBUIÇÃO"]["G3"] = "invalid historical time"
    workbook.save(path)  # This is explicitly generated test data, never the external XLSX.
    importer = LegacyImporter(path)
    with pytest.raises(ValidationError):
        importer.run()
    assert not Submission.objects.exists()
    assert LegacyImportRun.objects.get().status == "FAILED"
    assert LegacyImportIssue.objects.filter(code="INVALID_TIMESTAMP", row=3).exists()


def test_lenient_skips_timestamp_does_not_fabricate(tmp_path):
    path = source(tmp_path)
    workbook = openpyxl.load_workbook(path)
    workbook["DISTRIBUIÇÃO"]["F3"] = None
    workbook.save(path)
    importer = LegacyImporter(path, strict=False)
    importer.run()
    assert Submission.objects.count() == 1
    assert Submission.objects.get().processo_sei.endswith("000002/2025-01")
    assert LegacyImportIssue.objects.filter(code="INVALID_TIMESTAMP").exists()


def test_unknown_review_and_repeated_diligences_preserved(tmp_path):
    path = source(tmp_path)
    workbook = openpyxl.load_workbook(path)
    sei = workbook["DISTRIBUIÇÃO"]["E3"].value
    workbook["REVISÃO"]["C4"] = sei
    workbook["REVISÃO"]["CG4"] = "Sem docts"
    for row in (4, 5):
        workbook["DILIGÊNCIA"][f"C{row}"] = sei
        workbook["DILIGÊNCIA"][f"H{row}"] = f"raw event {row}"
    workbook.save(path)
    importer = LegacyImporter(path, strict=False)
    importer.run()
    sub = Submission.objects.get(processo_sei=sei)
    assert sub.reviews.get().preliminary_result == "PENDING_DECISION"
    assert LegacyImportIssue.objects.filter(
        code="UNKNOWN_REVIEW_RESULT", raw_value="Sem docts"
    ).exists()
    assert sub.diligences.count() == 2
    assert all(
        d.deadline is None
        and d.requested_at is None
        and d.requested_by is None
        and d.status == "LEGACY_UNKNOWN"
        for d in sub.diligences.all()
    )
    LegacyImporter(path, strict=False).run()
    assert sub.diligences.count() == 2
